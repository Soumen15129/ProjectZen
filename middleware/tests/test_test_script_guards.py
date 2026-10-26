"""
test_test_script_guards.py — the three guards behind the v16 "Invalid character /
found in sheet title" failure, with no LLM calls and no real database.

    python tests/test_test_script_guards.py

  1. The basic renderer cleans sheet titles Excel forbids; valid titles are unchanged.
  2. A Test Script reference is always treated as one register ('rows'), never as a
     tab per module.
  3. For Test Script the scope-gate repair may only add rows to existing sheets.
     Every other template keeps today's behaviour, including new sheets.
"""
import os, sys, json, types, asyncio, tempfile, shutil, contextlib, io

HERE = os.path.dirname(os.path.abspath(__file__))
MIDDLEWARE = os.path.dirname(HERE)
sys.path.insert(0, MIDDLEWARE)
TMP = tempfile.mkdtemp(prefix="pz-guards-test-")
os.environ["PROJECTZEN_DATA_DIR"] = TMP            # must precede any import of paths

# The module list that turned the v16 run into a tab-per-module workbook.
LONG_MODULES = [
    "Compensation - Annual Salary Review (merit, promotion, discretionary adjustment)",
    "Variable Pay - Annual Bonus with Global/Corporate and Individual sections",
    "Employee Central (EC) integration - data import, pay matrix/salary ranges",
    "Compensation / Total Rewards Statements (incl. local-language statements)",
]
CAPTURED = []


def _kind(p):
    if "Extract all project facts" in p:
        return "agent2"
    if "Return ONLY this JSON" in p:
        return "review"
    if "SCOPE CORRECTION" in p:
        return "repair"
    if "HARD RULES" in p:
        return "generate"
    return "text"


_CTX = json.dumps({"project_name": "Synthetic", "in_scope_modules": LONG_MODULES})
_PLAN = json.dumps({"title": "TS", "sheets": [
    {"name": "Test Scenarios", "headers": ["Test Scenario ID", "Scenario Description / Objective"],
     "rows": [["TS-SIT-CMP-001", "Prove that a planner can plan merit."]]},
    {"name": "Test Case Register", "headers": ["Test Case ID", "Test Level", "Country"],
     "rows": [["SIT-CMP-001-IN", "SIT", "India"]]}]})
_REPAIR = json.dumps({"additions": [
    {"sheet": "Test Case Register", "rows": [["SIT-CMP-002-IN", "SIT", "India"]]},
    {"sheet": "Compensation / Total Rewards", "headers": ["Test Case ID", "Test Steps"],
     "rows": [["SIT-STMT-001-IN", "1. Open statements"]]}]})


async def complete(prompt, system=None, model=None, on_chunk=None, **kw):
    kind = _kind(prompt)
    CAPTURED.append({"kind": kind, "prompt": prompt})
    out = {"agent2": _CTX, "review": '{"findings": []}', "generate": _PLAN,
           "repair": _REPAIR}.get(kind, "## 1. DELIVERABLE\nx\n")
    if on_chunk:
        await on_chunk(out)
    return out


class Conversation:
    session_id = None
    model = None

    def __init__(self, model=None, **kw):
        self.model = model

    async def __aenter__(self):
        return self

    async def __aexit__(self, *a):
        return None

    async def turn(self, prompt, model=None, on_chunk=None):
        return await complete(prompt, model=model, on_chunk=on_chunk)


@contextlib.contextmanager
def usage_scope(*a, **k):
    yield


async def _noop(*a, **k):
    return ""


stub = types.ModuleType("llm_client")
stub.complete, stub.Conversation, stub.usage_scope = complete, Conversation, usage_scope
stub.usage_summary = lambda *a, **k: {}
stub.LlmUnavailable = type("LlmUnavailable", (Exception,), {})
stub.FALLBACK_MODEL = "claude-sonnet-4-6"
stub.ToolSpec = lambda *a, **k: None
stub.run_tool_agent = stub.author_file = stub.preflight = _noop
sys.modules["llm_client"] = stub

import openpyxl
import templates
import scope_gate as sg
import adhoc_pipeline

FAILS = []


def check(label, cond, detail=""):
    print(f"  {'PASS' if cond else 'FAIL'}  {label}" + (f"   [{detail}]" if detail and not cond else ""))
    if not cond:
        FAILS.append(label)


def reference_like(path):
    """Sheet names that the long module list matched in v16, plus enough rows to
    make the depth gate ask for a repair."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Review Findings"
    ws.append(["#", "Finding", "Impact"])
    for name in ("Country Localisation Matrix", "Test Scenarios", "Test Case Register", "Test Steps"):
        ws = wb.create_sheet(name)
        ws.append(["Test Case ID", "Test Level", "Country"])
        for i in range(30):
            ws.append([f"SIT-EC-{i:03d}-IN", "SIT", "India"])
    wb.save(path)


async def run_pipeline(node_id, grounding_path):
    CAPTURED.clear()
    with contextlib.redirect_stdout(io.StringIO()) as out:
        res = await adhoc_pipeline.run(
            template="Test Script", node_id=node_id, output_format="xlsx",
            user_prompt="Template: Test Script\n\nGenerate the document.",
            input_text="REQ-1 Compensation worksheet launch for India.",
            grounding_path=grounding_path)
    return (res[0] if res else {}), out.getvalue()


async def main():
    print("\n1. Renderer: sheet titles")
    names = ["Compensation / Total Rewards Statements", "Q?:A*[B]\\C", "Test Steps",
             "Test Steps", "Roles & Test Users", "Country Localisation Matrix",
             "Variable Pay - LTI (PAT & RFO)", "Integration & Data-Load Inventory"]
    plan = {"title": "x", "sheets": [{"name": n, "headers": ["A", "B"], "rows": [["1", "two"]]}
                                     for n in names]}
    out = os.path.join(TMP, "render.xlsx")
    err = ""
    try:
        with contextlib.redirect_stdout(io.StringIO()):
            templates.generate_xlsx(plan, out)
    except Exception as e:
        err = f"{type(e).__name__}: {e}"
    check("a '/' or other forbidden character no longer crashes the renderer", not err, err)
    titles = openpyxl.load_workbook(out).sheetnames if not err else []
    check("forbidden characters replaced by '-'",
          titles[:2] == ["Compensation - Total Rewards St", "Q--A--B--C"], titles[:2])
    valid = [n for n in names if not any(c in n for c in "/\\?*[]:")]
    check("valid titles come back exactly as before",
          all(templates._sheet_title(n) == templates._safe(n)[:31] for n in valid))
    check("duplicate title still renamed by openpyxl, as before", titles[2:4] == ["Test Steps", "Test Steps1"],
          titles[2:4])
    check("every sheet written", len(titles) == len(names))

    print("\n2. Repair instruction")
    report = {"absent": ["Executive Review"], "in_rows_only": [("Statements", "Test Scenarios")]}
    default = sg.repair_instruction(report, row_shortfall=10, row_floor=50, row_actual=40)
    fixed = sg.repair_instruction(report, row_shortfall=10, row_floor=50, row_actual=40,
                                  allow_new_sheets=False)
    check("default: still offers a NEW sheet and a dedicated sheet per missing module",
          '"<name of a NEW sheet>"' in default and "Add a dedicated sheet for each" in default
          and "promote each to a dedicated sheet" in default)
    check("Test Script: no new sheet offered anywhere",
          "NEW sheet" not in fixed and "dedicated sheet" not in fixed and "Do NOT create new sheets" in fixed)
    check("Test Script: missing module asked for as rows in existing sheets",
          "Add rows for each to the EXISTING sheets" in fixed and "Executive Review" in fixed)
    check("both keep the depth and return-format rules",
          all("DEPTH SHORTFALL" in t and "RETURN ONLY THE ADDITIONS" in t for t in (default, fixed)))

    print("\n3. Repair merge")
    base = {"sheets": [{"name": "Test Case Register", "headers": ["Test Case ID", "Level"],
                        "rows": [["SIT-CMP-001-IN", "SIT"]]},
                       {"name": "Version History", "headers": ["Version"], "rows": [["v1"]]}]}
    delta = {"additions": [
        {"sheet": "Test Case Register", "rows": [["SIT-CMP-002-IN", "SIT"]]},
        {"sheet": "Annual Salary Review", "headers": ["Test Case ID", "Steps"], "rows": [["X-1", "y"]]}]}
    p1, n1, notes1 = sg.merge_row_additions(base, delta)
    check("default: new sheet created before Version History, as before",
          [s["name"] for s in p1["sheets"]] == ["Test Case Register", "Annual Salary Review",
                                                "Version History"] and n1 == 2, notes1)
    p2, n2, notes2 = sg.merge_row_additions(base, delta, allow_new_sheets=False)
    check("Test Script: new sheet skipped, existing sheet still gets its row",
          [s["name"] for s in p2["sheets"]] == ["Test Case Register", "Version History"] and n2 == 1
          and p2["sheets"][0]["rows"][-1] == ["SIT-CMP-002-IN", "SIT"], [s["name"] for s in p2["sheets"]])
    check("the skip is reported", any("skipped new sheet 'Annual Salary Review'" in n for n in notes2), notes2)
    check("input plan never mutated", [s["name"] for s in base["sheets"]] == ["Test Case Register",
                                                                             "Version History"]
          and len(base["sheets"][0]["rows"]) == 1)
    only_new = {"additions": [delta["additions"][1]]}
    check("Test Script: a delta of only new sheets changes nothing",
          sg.merge_row_additions(base, only_new, allow_new_sheets=False) == (base, 0, []))

    print("\n4. Whole adhoc pipeline, stub model")
    ref = os.path.join(TMP, "ref.xlsx")
    reference_like(ref)
    check("the v16 module list still reads as 'tabs' to the generic check (the misfire)",
          sg.reference_shape(ref, LONG_MODULES) == "tabs", sg.reference_shape(ref, LONG_MODULES))

    plan, log = await run_pipeline("test-script", ref)
    gen = [c["prompt"] for c in CAPTURED if c["kind"] == "generate"]
    rep = [c["prompt"] for c in CAPTURED if c["kind"] == "repair"]
    check("Test Script: treated as one register", "reference organises by: rows" in log)
    check("Test Script: generator is NOT told to make a tab per module",
          bool(gen) and "one dedicated sheet per module" not in gen[0]
          and "Do NOT impose a sheet-per-module layout" in gen[0])
    check("Test Script: repair pass ran and was told not to create sheets",
          bool(rep) and "Do NOT create new sheets" in rep[0] and "NEW sheet" not in rep[0])
    names = [s.get("name") for s in plan.get("sheets", [])]
    check("Test Script: no extra tab in the result", names == ["Test Scenarios", "Test Case Register"], names)
    reg = next((s for s in plan.get("sheets", []) if s.get("name") == "Test Case Register"), {})
    check("Test Script: the repair's row for an existing sheet was kept",
          ["SIT-CMP-002-IN", "SIT", "India"] in (reg.get("rows") or []), reg.get("rows"))

    plan, log = await run_pipeline("gov-matrix", ref)
    gen = [c["prompt"] for c in CAPTURED if c["kind"] == "generate"]
    names = [s.get("name") for s in plan.get("sheets", [])]
    check("other template: shape still measured from the reference ('tabs' here)",
          "reference organises by: tabs" in log and bool(gen) and "one dedicated sheet per module" in gen[0])
    check("other template: repair may still add a new sheet, as before",
          "Compensation / Total Rewards" in names, names)


if __name__ == "__main__":
    try:
        asyncio.run(main())
    finally:
        shutil.rmtree(TMP, ignore_errors=True)
    print("\nRESULT:", "PASS" if not FAILS else f"FAIL ({len(FAILS)}): {FAILS}")
    sys.exit(1 if FAILS else 0)
