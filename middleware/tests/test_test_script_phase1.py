"""
test_test_script_phase1.py — long cells, reference skeleton and rollup counting for
adhoc Test Script, with no LLM calls and no real database.

    python tests/test_test_script_phase1.py
"""
import os, sys, json, types, asyncio, tempfile, shutil, contextlib, io

HERE = os.path.dirname(os.path.abspath(__file__))
MIDDLEWARE = os.path.dirname(HERE)
sys.path.insert(0, MIDDLEWARE)
TMP = tempfile.mkdtemp(prefix="pz-phase1-test-")
os.environ["PROJECTZEN_DATA_DIR"] = TMP            # must precede any import of paths

LONG = "L" * 600
CAPTURED = []


def _kind(p):
    if "Extract all project facts" in p:
        return "agent2"
    if "Return ONLY this JSON" in p:
        return "review"
    if "HARD RULES" in p:
        return "generate"
    return "text"


_CTX = json.dumps({"project_name": "Synthetic", "in_scope_modules": ["Compensation"]})
_XLSX = json.dumps({"title": "TS", "sheets": [
    {"name": "Test Scenarios", "headers": ["Test Scenario ID", "Scenario Description / Objective"],
     "rows": [["TS-SIT-CMP-001", LONG]]},
    {"name": "Test Case Register", "headers": ["Test Case ID", "Test Level", "Country"],
     "rows": [["SIT-CMP-001-IN", "SIT", "India"]]}]})


async def complete(prompt, system=None, model=None, on_chunk=None, **kw):
    kind = _kind(prompt)
    CAPTURED.append({"kind": kind, "prompt": prompt})
    out = {"agent2": _CTX, "review": '{"findings": []}', "generate": _XLSX}.get(
        kind, "## 1. DELIVERABLE\nx\n")
    if on_chunk:
        await on_chunk(out)
    return out


class Conversation:
    session_id = None
    model = None

    def __init__(self, model=None, **kw):
        self.model = model

    async def __aenter__(self):
        return self

    async def __aexit__(self, *a):
        return None

    async def turn(self, prompt, model=None, on_chunk=None):
        return await complete(prompt, model=model, on_chunk=on_chunk)


@contextlib.contextmanager
def usage_scope(*a, **k):
    yield


async def _noop(*a, **k):
    return ""


stub = types.ModuleType("llm_client")
stub.complete, stub.Conversation, stub.usage_scope = complete, Conversation, usage_scope
stub.usage_summary = lambda *a, **k: {}
stub.LlmUnavailable = type("LlmUnavailable", (Exception,), {})
stub.FALLBACK_MODEL = "claude-sonnet-4-6"
stub.ToolSpec = lambda *a, **k: None
stub.run_tool_agent = stub.author_file = stub.preflight = _noop
sys.modules["llm_client"] = stub

import openpyxl
import input_budget
import cascade_agent as ca
import adhoc_pipeline
import reference_skeleton as rs
import scope_gate as sg
from quality.stages import rework

FAILS = []


def check(label, cond, detail=""):
    print(f"  {'PASS' if cond else 'FAIL'}  {label}" + (f"   [{detail}]" if detail and not cond else ""))
    if not cond:
        FAILS.append(label)


def gen_prompt():
    node = {"label": "Test Script", "tier": "T2", "phase": "", "type": "Document",
            "owner": "x", "trigger": "x", "notes": ""}
    return ca._build_gen_prompt("test-script", node, {"p": 1}, {}, "xlsx", "")


def rework_prompts():
    node = {"label": "Test Script"}
    f = [{"severity": "MAJOR", "section": "s", "issue": "i", "fix": "f"}]
    return (rework.build_prompt(node, '{"a": 1}', f, "{}"),
            rework.build_patch_prompt(node, '{"a": 1}', f, {"a": 1}))


def limits_seen():
    full, patch = rework_prompts()
    return (len(ca._safe_str("x" * 5000)),
            [n for n in (120, 1000) if f"max {n} chars" in gen_prompt()],
            [n for n in (120, 1000) if f"max {n} chars" in full],
            [n for n in (120, 1000) if f"max {n} chars" in patch])


def fixture_workbook(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cover"
    for r in (["Test Script v2"], ["Item", "Detail"], ["Document Title", "Reference Test Script"]):
        ws.append(r)
    ws = wb.create_sheet("Review Findings")
    ws.append(["#", "Round", "Category", "Finding"])
    ws.append(["R1-01", "R1", "Structure", "HISTORY-ONLY finding"])
    ws = wb.create_sheet("Test Scenarios")
    ws.append(["Test Scenario ID", "Scenario Key", "Test Scenario Name",
               "Scenario Description / Objective", "Test Level", "Test Case IDs"])
    for i in range(1, 6):
        ws.append([f"TS-SIT-EC-00{i}", f"HIRE-0{i}", f"Scenario {i}",
                   ("Prove that " + "detail " * 50) if i == 2 else f"Short {i}", "SIT",
                   f"SIT-EC-00{i}-IN"])
    ws = wb.create_sheet("Test Steps")
    ws.append(["Test Case ID", "Test Scenario ID", "Step No", "Step ID",
               "Action / Instruction", "Expected Result"])
    for i in range(1, 3):
        ws.append(["SIT-EC-001-IN", "TS-SIT-EC-001", i, f"SIT-EC-001-IN-S0{i}", f"A{i}", f"E{i}"])
    for i in range(1, 6):
        ws.append(["SIT-EC-002-IN", "TS-SIT-EC-002", i, f"SIT-EC-002-IN-S0{i}",
                   f"Open screen {i}", f"Screen {i} shows"])
    ws = wb.create_sheet("v1 Disposition")
    ws.append(["v1 Test Case ID", "v1 Sheet", "Disposition", "v2 Test Case ID(s)"])
    ws.append(["SIT-EC-001", "Old", "Superseded", "HISTORY-ONLY row"])
    ws = wb.create_sheet("Version History")
    ws.append(["Version", "Date", "Author", "Description"])
    ws.append(["v1", "2026-08-10", "Team", "Initial"])
    wb.save(path)


async def run_adhoc(node_id, grounding_path):
    CAPTURED.clear()
    with contextlib.redirect_stdout(io.StringIO()):
        res = await adhoc_pipeline.run(
            template="Test Script", node_id=node_id, output_format="xlsx",
            user_prompt="Template: Test Script\n\nGenerate the document.",
            input_text="REQ-1 Compensation worksheet launch for India.",
            grounding_text="FIRST-60K-OF-REFERENCE " * 20, grounding_path=grounding_path)
    return res


async def main():
    print("\n1. Which runs get long cells")
    b = adhoc_pipeline.long_cell_budget("test-script", "xlsx")
    check("Test Script xlsx -> 1000", b is not None and b.cell_chars == 1000)
    check("Test Script docx -> unchanged", adhoc_pipeline.long_cell_budget("test-script", "docx") is None)
    check("other template xlsx -> unchanged", adhoc_pipeline.long_cell_budget("gov-matrix", "xlsx") is None)
    check("no node id -> unchanged", adhoc_pipeline.long_cell_budget("", "xlsx") is None)
    check("every other budget field equals ADHOC",
          b is not None and {k: v for k, v in vars(b).items() if k != "cell_chars"}
          == {k: v for k, v in vars(input_budget.ADHOC).items() if k != "cell_chars"})

    print("\n2. Cell limit in shared code")
    check("cascade (no scope): cut 120, rules say 120", limits_seen() == (120, [120], [120], [120]),
          limits_seen())
    with input_budget.scope(input_budget.ADHOC):
        check("other adhoc templates: still 120", limits_seen() == (120, [120], [120], [120]),
              limits_seen())
    with input_budget.scope(b):
        check("Test Script: cut 1000, rules say 1000", limits_seen() == (1000, [1000], [1000], [1000]),
              limits_seen())
    check("no leak after the scope closes", limits_seen()[0] == 120)

    async def in_scope():
        with input_budget.scope(b):
            await asyncio.sleep(0.01)
            return len(ca._safe_str("x" * 5000))

    async def outside():
        await asyncio.sleep(0.005)
        return len(ca._safe_str("x" * 5000))

    a, c = await asyncio.gather(in_scope(), outside())
    check("concurrent cascade task keeps 120 while Test Script runs", (a, c) == (1000, 120), (a, c))

    print("\n3. Reference skeleton")
    path = os.path.join(TMP, "ref.xlsx")
    fixture_workbook(path)
    s = rs.build(path)
    check("built", bool(s))
    check("history sheets omitted", 'SHEET "Review Findings"' not in s and 'SHEET "v1 Disposition"' not in s
          and "HISTORY-ONLY" not in s)
    check("real sheets kept, in order",
          [ln.split('"')[1] for ln in s.splitlines() if ln.startswith("SHEET")]
          == ["Cover", "Test Scenarios", "Test Steps", "Version History"])
    check("exact headers listed",
          "COLUMNS (6, use exactly): Test Scenario ID | Scenario Key | Test Scenario Name | "
          "Scenario Description / Objective | Test Level | Test Case IDs" in s)
    check("richest scenario row shown in full", ("Prove that " + "detail " * 50).strip()[:300] in s)
    steps = s.split('SHEET "Test Steps"')[1].split("SHEET")[0]
    check("steps example is one whole case (5 consecutive steps)",
          steps.count("SIT-EC-002-IN-S0") == 5 and "SIT-EC-001-IN-S0" not in steps)
    check("Cover shown as rows", "Document Title | Reference Test Script" in s)
    check("deterministic", rs.build(path) == s)
    check("size bound respected", len(rs.build(path, max_chars=900)) <= 900)
    check("non-workbook -> empty", rs.build(os.path.join(TMP, "x.docx")) == "")
    check("missing file -> empty", rs.build(os.path.join(TMP, "missing.xlsx")) == "")
    real = os.path.join(os.path.dirname(MIDDLEWARE), "storage", "document_store", "refs",
                        "d9397f7e-d0f5-4783-a239-aad489c7dc35.xlsx")
    if os.path.exists(real):
        r = rs.build(real)
        check(f"real Test Script reference: {len(r):,} chars, within 30,000", 0 < len(r) <= 30_000)
        check("real reference: all three core sheets with their columns",
              all(f'SHEET "{n}"' in r for n in ("Test Scenarios", "Test Case Register", "Test Steps"))
              and "Scenario Description / Objective" in r)
        check("real reference: history sheets omitted",
              'SHEET "Review Findings"' not in r and 'SHEET "v1 Disposition"' not in r)

    hint = ("REFERENCE ORGANISATION — measured:\n  It divides its content into 4 parts:\n"
            "    - Review Findings: 55 rows (2% of its content)\n"
            "    - Test Steps: 1766 rows (54% of its content)\n"
            "    - v1 Disposition: 385 rows (12% of its content)\n"
            "    - RTM: 201 rows (6% of its content)\n  THIS SPLIT IS REAL")
    h = rs.drop_history_parts(hint)
    check("organisation hint: history parts removed, count corrected",
          "Review Findings" not in h and "v1 Disposition" not in h and "into 2 parts" in h
          and "Test Steps: 1766" in h)
    check("organisation hint without history -> unchanged",
          rs.drop_history_parts(h) == h and rs.drop_history_parts("") == "")

    print("\n4. Dashboard recount")
    dash = lambda labels: {"name": "Dashboard", "headers": ["Test Level", "Total TCs"],
                           "rows": [[x, ""] for x in labels]}
    plan = {"sheets": [
        dash(["SIT", "UAT", "Total"]),
        {"name": "Test Scenarios", "headers": ["Test Scenario ID", "Test Case IDs", "Test Level"],
         "rows": [["TS-SIT-CMP-001", "SIT-CMP-001-IN; SIT-CMP-001-NP", "SIT"],
                  ["TS-UAT-CMP-002", "UAT-CMP-002-IN", "UAT"]]},
        {"name": "Test Case Register", "headers": ["Test Case ID", "Test Level", "Country"],
         "rows": [["SIT-CMP-001-IN", "SIT", "India"], ["SIT-CMP-001-NP", "SIT", "Nepal"],
                  ["UAT-CMP-002-IN", "UAT", "India"]]},
        {"name": "Test Steps", "headers": ["Test Case ID", "Step No", "Test Level", "Country"],
         "rows": [[cid, str(n), lv, co] for cid, lv, co in
                  (("SIT-CMP-001-IN", "SIT", "India"), ("SIT-CMP-001-NP", "SIT", "Nepal"),
                   ("UAT-CMP-002-IN", "UAT", "India")) for n in range(1, 5)]}]}
    out, _ = sg.recompute_rollup(plan)
    got = {r[0]: r[1] for r in out["sheets"][0]["rows"]}
    check("grounding layout: 3 cases, not scenarios or steps", got == {"SIT": "2", "UAT": "1", "Total": "3"}, got)

    old = {"sheets": [
        {"name": "Test Summary", "headers": ["Category", "Total TCs"],
         "rows": [["By Test Level - SIT", ""], ["By Test Level - UAT", ""],
                  ["By Country - India", ""], ["Total", ""]]},
        {"name": "SIT Test Cases", "headers": ["Test Case ID", "Test Level", "Country"],
         "rows": [[f"SIT-EC-00{i}", "SIT", "India"] for i in range(1, 6)]},
        {"name": "UAT Test Cases", "headers": ["Test Case ID", "Test Level", "Country"],
         "rows": [[f"UAT-EC-00{i}", "UAT", "Nepal"] for i in range(1, 4)]}]}
    out, _ = sg.recompute_rollup(old)
    got = {r[0]: r[1] for r in out["sheets"][0]["rows"]}
    check("earlier layout: counts unchanged (SIT 5, UAT 3, India 5, Total 8)",
          got == {"By Test Level - SIT": "5", "By Test Level - UAT": "3",
                  "By Country - India": "5", "Total": "8"}, got)
    none = {"sheets": [{"name": "Plan", "headers": ["Task", "Owner"], "rows": [["a", "b"]]}]}
    check("no 'Total TCs' column -> untouched", sg.recompute_rollup(none) == (none, []))
    check("ID recognition", [sg._is_case_id(v) for v in
                             ("SIT-EC-001-IN", "SIT-DM-001", "CMP-IN-001", "EC-Core", "1", "")]
          == [True, True, True, False, False, False])

    print("\n5. Whole adhoc pipeline, stub model")
    res = await run_adhoc("test-script", path)
    plan_out = res[0] if res else {}
    cell = plan_out.get("sheets", [{}])[0].get("rows", [[None, ""]])[0][1] if plan_out else ""
    check("Test Script: 600-char objective survives intact", len(cell) == 600, len(cell))
    gen = [c["prompt"] for c in CAPTURED if c["kind"] == "generate"]
    check("generate prompt allows 1000-char cells", bool(gen) and "max 1000 chars" in gen[0]
          and "max 120 chars" not in gen[0])
    check("every stage that reads the reference sees the skeleton, not the cut text",
          any("REFERENCE WORKBOOK SKELETON" in c["prompt"] for c in CAPTURED if c["kind"] == "text")
          and bool(gen) and "REFERENCE WORKBOOK SKELETON" in gen[0]
          and not any("FIRST-60K-OF-REFERENCE" in c["prompt"] for c in CAPTURED))
    check("no prompt carries the reference's history sheets",
          not any("HISTORY-ONLY" in c["prompt"] for c in CAPTURED))
    check("no leak after the run", len(ca._safe_str("x" * 5000)) == 120)

    res = await run_adhoc("gov-matrix", path)
    plan_out = res[0] if res else {}
    cell = plan_out.get("sheets", [{}])[0].get("rows", [[None, ""]])[0][1] if plan_out else ""
    gen = [c["prompt"] for c in CAPTURED if c["kind"] == "generate"]
    check("other template: cells still cut at 120", len(cell) == 120, len(cell))
    check("other template: rule still 120, reference text as before",
          bool(gen) and "max 120 chars" in gen[0] and "FIRST-60K-OF-REFERENCE" in gen[0]
          and not any("REFERENCE WORKBOOK SKELETON" in c["prompt"] for c in CAPTURED))


if __name__ == "__main__":
    try:
        asyncio.run(main())
    finally:
        shutil.rmtree(TMP, ignore_errors=True)
    print("\nRESULT:", "PASS" if not FAILS else f"FAIL ({len(FAILS)}): {FAILS}")
    sys.exit(1 if FAILS else 0)
